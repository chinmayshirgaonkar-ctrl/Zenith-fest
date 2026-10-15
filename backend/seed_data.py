import os
from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db, is_postgres

def seed_events(seed_path: str = None):
    """
    Seeds the events table from the companion Excel file at data/events_seed.xlsx.
    Raises FileNotFoundError if data/events_seed.xlsx is missing.
    Skips seeding if the events table already has rows.
    """
    # Ensure tables are created first
    init_db()

    if seed_path is None:
        seed_path = Path(__file__).resolve().parent.parent / "data" / "events_seed.xlsx"
    else:
        seed_path = Path(seed_path)

    # Check if seed spreadsheet exists
    if not seed_path.exists():
        raise FileNotFoundError(
            f"Seed data file not found at '{seed_path}'. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events table already contains data
    cursor.execute("SELECT COUNT(*) AS cnt FROM events")
    count = cursor.fetchone()["cnt"]
    if count > 0:
        print(f"Events table already has {count} rows. Skipping seed.")
        conn.close()
        return

    # Load workbook and sheet
    wb = openpyxl.load_workbook(str(seed_path), data_only=True)
    if "Events" not in wb.sheetnames:
        raise ValueError("Sheet 'Events' not found in " + str(seed_path))
    sheet = wb["Events"]

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("Warning: 'Events' sheet is empty.")
        conn.close()
        return

    # Row 1 is header: id, title, category, description, event_date, event_time, venue, capacity, fee, icon
    header = [str(col).strip().lower() if col is not None else "" for col in rows[0]]
    expected_header = ["id", "title", "category", "description", "event_date", "event_time", "venue", "capacity", "fee", "icon"]
    
    # Insert every subsequent row
    inserted_count = 0
    for row in rows[1:]:
        if not row or row[0] is None or row[1] is None:
            continue
        
        event_id = int(row[0])
        title = str(row[1]).strip()
        category = str(row[2]).strip()
        description = str(row[3]).strip() if row[3] is not None else ""
        event_date = str(row[4]).strip()
        event_time = str(row[5]).strip()
        venue = str(row[6]).strip()
        capacity = int(row[7])
        fee = int(row[8]) if row[8] is not None else 0
        icon = str(row[9]).strip() if row[9] is not None else "📅"

        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon))
        inserted_count += 1

    if is_postgres():
        try:
            cursor.execute("SELECT setval(pg_get_serial_sequence('events', 'id'), COALESCE((SELECT MAX(id) FROM events), 1))")
        except Exception as e:
            print(f"Notice on PostgreSQL sequence sync: {e}")

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {seed_path}.")

if __name__ == "__main__":
    import sys
    if sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
        sys.stdout.reconfigure(encoding="utf-8")
    seed_events()
    
    # Print out events to verify
    conn = get_db_connection()
    events = conn.execute("SELECT id, title, category, event_date, capacity, fee, icon FROM events ORDER BY id").fetchall()
    print(f"\nTotal rows in events table: {len(events)}")
    for ev in events:
        print(f"[{ev['id']}] {ev['icon']} {ev['title']} ({ev['category']}) - {ev['event_date']} - Cap: {ev['capacity']} - Fee: Rs.{ev['fee']}")
    conn.close()
